"""Best-effort document date extraction with quality rating.

Tries format-specific metadata first (PDF, EPUB, DOCX, XLSX), then falls
back to filesystem ``mtime``.  Each source is assigned a quality tier so
downstream consumers (LLM prompts, API responses) can communicate how
trustworthy the date is.

Quality tiers (highest → lowest):
    "document_metadata" — explicit creation/publish date inside the file format
    "filesystem"        — OS modification time (unreliable after copy/sync)
    "none"              — no date could be determined
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

logger = logging.getLogger(__name__)

QUALITY_DOCUMENT_METADATA = "document_metadata"
QUALITY_FILESYSTEM = "filesystem"
QUALITY_NONE = "none"


@dataclass
class DocumentDate:
    """Extracted document date with provenance quality."""

    date: str | None  # ISO-8601 date string (YYYY-MM-DD or YYYY-MM-DDThh:mm:ss+00:00)
    quality: str      # one of the QUALITY_* constants
    source: str       # human-readable origin, e.g. "pdf:creationDate", "filesystem:mtime"

    @property
    def date_brief(self) -> str:
        """Short date for display (YYYY-MM-DD or empty)."""
        if not self.date:
            return ""
        return self.date[:10]


_NONE = DocumentDate(date=None, quality=QUALITY_NONE, source="none")


# ---------------------------------------------------------------------------
# PDF  (PyMuPDF metadata dict)
# ---------------------------------------------------------------------------

_PDF_DATE_RE = re.compile(
    r"D:(\d{4})(\d{2})?(\d{2})?(\d{2})?(\d{2})?(\d{2})?"
)


def _parse_pdf_date(raw: str) -> str | None:
    """Parse PDF date string ``D:YYYYMMDDHHmmSS…`` → ISO-8601."""
    m = _PDF_DATE_RE.search(raw)
    if not m:
        return None
    year = m.group(1)
    month = m.group(2) or "01"
    day = m.group(3) or "01"
    hour = m.group(4) or "00"
    minute = m.group(5) or "00"
    second = m.group(6) or "00"
    try:
        dt = datetime(
            int(year), int(month), int(day),
            int(hour), int(minute), int(second),
            tzinfo=timezone.utc,
        )
        return dt.isoformat()
    except ValueError:
        return None


def extract_pdf_date(pdf_path: Path) -> DocumentDate:
    """Extract date from PDF metadata (creationDate preferred, then modDate)."""
    try:
        import fitz
        doc = fitz.open(pdf_path)
        meta = doc.metadata or {}
        doc.close()
    except Exception as exc:
        logger.debug("Cannot read PDF metadata from '%s': %s", pdf_path.name, exc)
        return _NONE

    for key in ("creationDate", "modDate"):
        raw = (meta.get(key) or "").strip()
        if raw:
            iso = _parse_pdf_date(raw)
            if iso:
                return DocumentDate(
                    date=iso,
                    quality=QUALITY_DOCUMENT_METADATA,
                    source=f"pdf:{key}",
                )
    return _NONE


# ---------------------------------------------------------------------------
# EPUB  (Dublin Core dc:date in OPF metadata)
# ---------------------------------------------------------------------------

def extract_epub_date(epub_path: Path) -> DocumentDate:
    """Extract date from EPUB Dublin Core metadata."""
    try:
        from ebooklib import epub  # type: ignore[import-untyped]
        book = epub.read_epub(str(epub_path), options={"ignore_ncx": True})
    except Exception as exc:
        logger.debug("Cannot read EPUB metadata from '%s': %s", epub_path.name, exc)
        return _NONE

    try:
        dates = book.get_metadata("DC", "date")
    except Exception:
        return _NONE

    for entry in dates:
        raw = (entry[0] if isinstance(entry, tuple) else str(entry)).strip()
        if raw:
            iso = _normalise_date_string(raw)
            if iso:
                return DocumentDate(
                    date=iso,
                    quality=QUALITY_DOCUMENT_METADATA,
                    source="epub:dc:date",
                )
    return _NONE


# ---------------------------------------------------------------------------
# DOCX  (core_properties.created / modified)
# ---------------------------------------------------------------------------

def extract_docx_date(docx_path: Path) -> DocumentDate:
    """Extract date from DOCX core properties."""
    try:
        from docx import Document  # type: ignore[import-untyped]
        doc = Document(str(docx_path))
        props = doc.core_properties
    except Exception as exc:
        logger.debug("Cannot read DOCX metadata from '%s': %s", docx_path.name, exc)
        return _NONE

    for attr, label in [("created", "docx:created"), ("modified", "docx:modified")]:
        dt = getattr(props, attr, None)
        if isinstance(dt, datetime):
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            return DocumentDate(
                date=dt.isoformat(),
                quality=QUALITY_DOCUMENT_METADATA,
                source=label,
            )
    return _NONE


# ---------------------------------------------------------------------------
# XLSX  (openpyxl workbook properties)
# ---------------------------------------------------------------------------

def extract_xlsx_date(xlsx_path: Path) -> DocumentDate:
    """Extract date from XLSX workbook properties."""
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]
        wb = load_workbook(str(xlsx_path), read_only=True, data_only=True)
        props = wb.properties
    except Exception as exc:
        logger.debug("Cannot read XLSX metadata from '%s': %s", xlsx_path.name, exc)
        return _NONE

    for attr, label in [("created", "xlsx:created"), ("modified", "xlsx:modified")]:
        dt = getattr(props, attr, None)
        if isinstance(dt, datetime):
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            try:
                wb.close()
            except Exception:
                pass
            return DocumentDate(
                date=dt.isoformat(),
                quality=QUALITY_DOCUMENT_METADATA,
                source=label,
            )
    try:
        wb.close()
    except Exception:
        pass
    return _NONE


# ---------------------------------------------------------------------------
# Filesystem fallback
# ---------------------------------------------------------------------------

def extract_filesystem_date(path: Path) -> DocumentDate:
    """Use file modification time as a last-resort date."""
    try:
        mtime = path.stat().st_mtime
        dt = datetime.fromtimestamp(mtime, tz=timezone.utc)
        return DocumentDate(
            date=dt.isoformat(),
            quality=QUALITY_FILESYSTEM,
            source="filesystem:mtime",
        )
    except Exception:
        return _NONE


# ---------------------------------------------------------------------------
# Unified dispatcher
# ---------------------------------------------------------------------------

_FORMAT_EXTRACTORS: dict[str, Callable[[Path], DocumentDate]] = {
    ".pdf": extract_pdf_date,
    ".epub": extract_epub_date,
    ".docx": extract_docx_date,
    ".xlsx": extract_xlsx_date,
    ".xls": extract_xlsx_date,
}


def extract_document_date(path: Path) -> DocumentDate:
    """Best-effort date extraction: format-specific metadata → filesystem mtime.

    Returns a ``DocumentDate`` with quality rating so callers know how
    trustworthy the date is.
    """
    suffix = path.suffix.lower()
    extractor = _FORMAT_EXTRACTORS.get(suffix)
    if extractor is not None:
        result = extractor(path)
        if result.date:
            logger.debug(
                "Date for '%s': %s (quality=%s, source=%s)",
                path.name, result.date_brief, result.quality, result.source,
            )
            return result

    fs = extract_filesystem_date(path)
    if fs.date:
        logger.debug(
            "Date for '%s': %s (filesystem fallback)", path.name, fs.date_brief,
        )
        return fs

    return _NONE


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _normalise_date_string(raw: str) -> str | None:
    """Try to parse a free-form date string into ISO-8601."""
    raw = raw.strip()
    for fmt in (
        "%Y-%m-%dT%H:%M:%S%z",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y",
    ):
        try:
            dt = datetime.strptime(raw, fmt)
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            return dt.isoformat()
        except ValueError:
            continue
    return None
